"""Tests for app.services.accrual_report_import — Ozon's own downloadable
«Начисления» XLSX report (Финансы → Начисления → «Скачать отчёт»),
aggregated here into one row per (accrual_date, service_group, charge_type)
on AccrualReportDailyStatistic. See that module's/model's own docstrings
for why this manual import exists (no confirmed API equivalent at this
granularity) and why it's a more precise source than CashFlowStatementPeriod."""
import datetime as dt
import io

import openpyxl
import pytest

from app.models.accrual_report_daily_statistic import AccrualReportDailyStatistic
from app.models.product_accrual_report_daily_statistic import ProductAccrualReportDailyStatistic
from app.services.accrual_report_import import import_accrual_report_from_file
from tests.conftest import login

_HEADER = [
    "ID начисления", "Дата начисления", "Группа услуг", "Тип начисления", "Артикул", "SKU",
    "Название товара", "Количество", "Цена продавца",
    "Дата принятия заказа в обработку или оказания услуги", "Платформа продажи", "Схема работы",
    "Вознаграждение Ozon, %", "Индекс локализации, %", "Среднее время доставки, часы", "Сумма итого, руб.",
]


def _build_xlsx(rows: list[dict], *, period: str = "Период: 01.09.2026-21.09.2026") -> bytes:
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.append([period])
    ws.append(_HEADER)
    for r in rows:
        ws.append([
            r.get("id", "id-1"), r["date"], r["group"], r["type"], r.get("offer_id", "art"),
            r.get("sku", "111"), r.get("name", "Товар"), r.get("qty", 1), r.get("seller_price", 0),
            r.get("date", ""), r.get("platform", ""), r.get("scheme", ""), 0, None, None, r["amount"],
        ])
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def test_aggregates_by_date_group_type(db_session, two_stores_with_users):
    d = two_stores_with_users
    store_id = d["store_a"].id

    content = _build_xlsx([
        {"date": dt.datetime(2026, 9, 1), "group": "Услуги партнёров", "type": "Эквайринг", "amount": -66.75},
        {"date": dt.datetime(2026, 9, 1), "group": "Услуги партнёров", "type": "Эквайринг", "amount": -31.56},
        {"date": dt.datetime(2026, 9, 2), "group": "Услуги партнёров", "type": "Эквайринг", "amount": -45.6},
    ])

    result = import_accrual_report_from_file(db_session, store_id=store_id, filename="acc.xlsx", content=content)
    db_session.commit()

    assert result.created == 2
    assert not result.errors

    rows = db_session.query(AccrualReportDailyStatistic).filter(AccrualReportDailyStatistic.store_id == store_id).all()
    by_date = {r.accrual_date: r for r in rows}
    assert len(rows) == 2
    row_sep1 = by_date[dt.date(2026, 9, 1)]
    assert row_sep1.service_group == "Услуги партнёров"
    assert row_sep1.charge_type == "Эквайринг"
    assert float(row_sep1.amount_rub) == pytest.approx(-98.31)
    assert row_sep1.rows_count == 2
    row_sep2 = by_date[dt.date(2026, 9, 2)]
    assert float(row_sep2.amount_rub) == pytest.approx(-45.6)
    assert row_sep2.rows_count == 1


def test_different_groups_and_types_kept_separate(db_session, two_stores_with_users):
    d = two_stores_with_users
    store_id = d["store_a"].id

    content = _build_xlsx([
        {"date": dt.datetime(2026, 9, 1), "group": "Услуги доставки", "type": "Логистика", "amount": -100.0},
        {"date": dt.datetime(2026, 9, 1), "group": "Услуги доставки", "type": "Обратная логистика", "amount": -20.0},
        {"date": dt.datetime(2026, 9, 1), "group": "Другие услуги и штрафы", "type": "Утилизация товара", "amount": -5.0},
    ])

    result = import_accrual_report_from_file(db_session, store_id=store_id, filename="acc.xlsx", content=content)
    db_session.commit()

    assert result.created == 3
    rows = db_session.query(AccrualReportDailyStatistic).filter(AccrualReportDailyStatistic.store_id == store_id).all()
    keys = {(r.service_group, r.charge_type): float(r.amount_rub) for r in rows}
    assert keys == {
        ("Услуги доставки", "Логистика"): -100.0,
        ("Услуги доставки", "Обратная логистика"): -20.0,
        ("Другие услуги и штрафы", "Утилизация товара"): -5.0,
    }


def test_reupload_overlapping_range_replaces_not_duplicates(db_session, two_stores_with_users):
    """Uploading a second file whose date range overlaps the first must
    REPLACE the sum for the days/groups/types it covers, never add on top
    — see the model's own docstring on why the unique key is (store_id,
    accrual_date, service_group, charge_type), not the uploaded file's own
    period range."""
    d = two_stores_with_users
    store_id = d["store_a"].id

    first = _build_xlsx([
        {"date": dt.datetime(2026, 9, 1), "group": "Услуги партнёров", "type": "Эквайринг", "amount": -50.0},
    ])
    import_accrual_report_from_file(db_session, store_id=store_id, filename="acc1.xlsx", content=first)
    db_session.commit()

    second = _build_xlsx([
        {"date": dt.datetime(2026, 9, 1), "group": "Услуги партнёров", "type": "Эквайринг", "amount": -70.0},
        {"date": dt.datetime(2026, 9, 1), "group": "Услуги партнёров", "type": "Эквайринг", "amount": -10.0},
    ])
    import_accrual_report_from_file(db_session, store_id=store_id, filename="acc2.xlsx", content=second)
    db_session.commit()

    rows = db_session.query(AccrualReportDailyStatistic).filter(AccrualReportDailyStatistic.store_id == store_id).all()
    assert len(rows) == 1
    assert float(rows[0].amount_rub) == pytest.approx(-80.0)
    assert rows[0].rows_count == 2


def test_per_sku_aggregation_alongside_store_level(db_session, two_stores_with_users):
    """The SAME rows must ALSO be aggregated per SKU into
    ProductAccrualReportDailyStatistic (requested 2026-09-22 so «РНП
    Товары» can show per-product logistics costs) — without changing the
    store-level totals at all."""
    d = two_stores_with_users
    store_id = d["store_a"].id

    content = _build_xlsx([
        {"date": dt.datetime(2026, 9, 1), "group": "Услуги партнёров", "type": "Эквайринг", "sku": "111", "amount": -30.0},
        {"date": dt.datetime(2026, 9, 1), "group": "Услуги партнёров", "type": "Эквайринг", "sku": "222", "amount": -20.0},
        {"date": dt.datetime(2026, 9, 2), "group": "Услуги доставки", "type": "Логистика", "sku": "111", "amount": -100.0},
    ])

    result = import_accrual_report_from_file(db_session, store_id=store_id, filename="acc.xlsx", content=content)
    db_session.commit()

    assert result.created == 2  # store-level: (09-01, Услуги партнёров, Эквайринг) + (09-02, Услуги доставки, Логистика)
    store_rows = db_session.query(AccrualReportDailyStatistic).filter(AccrualReportDailyStatistic.store_id == store_id).all()
    assert len(store_rows) == 2
    acquiring_row = next(r for r in store_rows if r.charge_type == "Эквайринг")
    assert float(acquiring_row.amount_rub) == pytest.approx(-50.0)  # 111 + 222 combined, unchanged by the per-SKU split

    sku_rows = {
        (r.ozon_sku, r.accrual_date, r.charge_type): float(r.amount_rub)
        for r in db_session.query(ProductAccrualReportDailyStatistic).filter(ProductAccrualReportDailyStatistic.store_id == store_id).all()
    }
    assert sku_rows == {
        ("111", dt.date(2026, 9, 1), "Эквайринг"): -30.0,
        ("222", dt.date(2026, 9, 1), "Эквайринг"): -20.0,
        ("111", dt.date(2026, 9, 2), "Логистика"): -100.0,
    }


def test_per_sku_reupload_overlapping_replaces_not_duplicates(db_session, two_stores_with_users):
    d = two_stores_with_users
    store_id = d["store_a"].id

    first = _build_xlsx([
        {"date": dt.datetime(2026, 9, 1), "group": "Услуги партнёров", "type": "Эквайринг", "sku": "111", "amount": -30.0},
    ])
    import_accrual_report_from_file(db_session, store_id=store_id, filename="acc1.xlsx", content=first)
    db_session.commit()

    second = _build_xlsx([
        {"date": dt.datetime(2026, 9, 1), "group": "Услуги партнёров", "type": "Эквайринг", "sku": "111", "amount": -45.0},
    ])
    import_accrual_report_from_file(db_session, store_id=store_id, filename="acc2.xlsx", content=second)
    db_session.commit()

    rows = db_session.query(ProductAccrualReportDailyStatistic).filter(ProductAccrualReportDailyStatistic.store_id == store_id).all()
    assert len(rows) == 1
    assert float(rows[0].amount_rub) == pytest.approx(-45.0)


def test_missing_sku_column_still_imports_store_level_only(db_session, two_stores_with_users):
    """A file without a SKU column (optional, unlike Группа услуг/Тип
    начисления) still imports the store-level totals — only the per-product
    table stays empty for that upload, see the module's own docstring."""
    d = two_stores_with_users
    store_id = d["store_a"].id

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.append(["Период: 01.09.2026-21.09.2026"])
    ws.append(["Дата начисления", "Группа услуг", "Тип начисления", "Сумма итого, руб."])
    ws.append([dt.datetime(2026, 9, 1), "Услуги партнёров", "Эквайринг", -30.0])
    buf = io.BytesIO()
    wb.save(buf)

    result = import_accrual_report_from_file(db_session, store_id=store_id, filename="acc.xlsx", content=buf.getvalue())
    db_session.commit()

    assert result.created == 1
    assert not result.errors
    store_rows = db_session.query(AccrualReportDailyStatistic).filter(AccrualReportDailyStatistic.store_id == store_id).all()
    assert len(store_rows) == 1
    sku_rows = db_session.query(ProductAccrualReportDailyStatistic).filter(ProductAccrualReportDailyStatistic.store_id == store_id).all()
    assert sku_rows == []


def test_missing_required_columns_reports_error(db_session, two_stores_with_users):
    d = two_stores_with_users
    store_id = d["store_a"].id

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.append(["Период: 01.09.2026-21.09.2026"])
    ws.append(["Группа услуг", "Тип начисления", "Что-то ещё"])
    ws.append(["Услуги партнёров", "Эквайринг", 5])
    buf = io.BytesIO()
    wb.save(buf)

    result = import_accrual_report_from_file(db_session, store_id=store_id, filename="acc.xlsx", content=buf.getvalue())

    assert result.created == 0
    assert any("отсутствуют обязательные колонки" in e for e in result.errors)


def test_non_xlsx_extension_rejected(db_session, two_stores_with_users):
    d = two_stores_with_users
    store_id = d["store_a"].id
    result = import_accrual_report_from_file(db_session, store_id=store_id, filename="acc.csv", content=b"whatever")
    assert result.created == 0
    assert any(".xlsx" in e for e in result.errors)


def test_unparseable_date_skipped_and_reported(db_session, two_stores_with_users):
    d = two_stores_with_users
    store_id = d["store_a"].id

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.append(["Период: 01.09.2026-21.09.2026"])
    ws.append(_HEADER)
    ws.append(["id-1", "не дата", "Услуги партнёров", "Эквайринг", "art", "111", "Товар", 1, 0, "", "", "", 0, None, None, -10.0])
    ws.append(["id-2", dt.datetime(2026, 9, 1), "Услуги партнёров", "Эквайринг", "art", "111", "Товар", 1, 0, "", "", "", 0, None, None, -5.0])
    buf = io.BytesIO()
    wb.save(buf)

    result = import_accrual_report_from_file(db_session, store_id=store_id, filename="acc.xlsx", content=buf.getvalue())
    db_session.commit()

    assert result.created == 1
    assert any("нераспознанной датой" in e for e in result.errors)
    rows = db_session.query(AccrualReportDailyStatistic).filter(AccrualReportDailyStatistic.store_id == store_id).all()
    assert len(rows) == 1
    assert float(rows[0].amount_rub) == pytest.approx(-5.0)


def test_upload_route_end_to_end(client, db_session, two_stores_with_users):
    d = two_stores_with_users
    store_id = d["store_a"].id
    content = _build_xlsx([
        {"date": dt.datetime(2026, 9, 1), "group": "Услуги партнёров", "type": "Эквайринг", "amount": -66.75},
    ])

    login(client, "owner_a@example.com", "password123")
    resp = client.post(
        f"/api/stores/{store_id}/dashboard/upload-accruals",
        files={"file": ("acc.xlsx", content, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")},
    )
    assert resp.status_code == 200
    body = resp.json()
    assert body["created"] == 1
    assert body["errors"] == []

    rows = db_session.query(AccrualReportDailyStatistic).filter(AccrualReportDailyStatistic.store_id == store_id).all()
    assert len(rows) == 1


def test_upload_route_requires_manager_role(client, db_session, two_stores_with_users):
    from app.core.security import hash_password
    from app.models.membership import StoreMembership, StoreRole
    from app.models.user import User

    d = two_stores_with_users
    viewer = User(email="acc_viewer@example.com", full_name="Viewer", password_hash=hash_password("password123"))
    db_session.add(viewer)
    db_session.flush()
    db_session.add(StoreMembership(user_id=viewer.id, store_id=d["store_a"].id, role=StoreRole.VIEWER))
    db_session.commit()

    content = _build_xlsx([{"date": dt.datetime(2026, 9, 1), "group": "Услуги партнёров", "type": "Эквайринг", "amount": -66.75}])
    login(client, "acc_viewer@example.com", "password123")
    resp = client.post(
        f"/api/stores/{d['store_a'].id}/dashboard/upload-accruals",
        files={"file": ("acc.xlsx", content, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")},
    )
    assert resp.status_code == 403
